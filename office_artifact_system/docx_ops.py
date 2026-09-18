from __future__ import annotations

from collections import Counter
from hashlib import sha256
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from .common import Finding, hash_file


def _paragraph_texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def _table_texts(doc: Document) -> list[list[list[str]]]:
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]


def content_signature(path: str | Path) -> dict[str, Any]:
    doc = Document(path)
    paragraphs = _paragraph_texts(doc)
    tables = _table_texts(doc)
    payload = repr((paragraphs, tables)).encode("utf-8")
    return {
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "text_sha256": sha256(payload).hexdigest(),
    }


def inspect_docx(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    doc = Document(path)
    styles = Counter(p.style.name if p.style else "<none>" for p in doc.paragraphs)
    fonts = Counter()
    sizes = Counter()
    manual_bullets = []
    for i, p in enumerate(doc.paragraphs, 1):
        if p.text.lstrip().startswith(("•", "", "●", "◦")):
            manual_bullets.append(i)
        for r in p.runs:
            if r.text:
                fonts[r.font.name or "<inherit>"] += len(r.text)
                if r.font.size:
                    sizes[round(r.font.size.pt, 2)] += len(r.text)
    sections = []
    for s in doc.sections:
        sections.append({
            "page_width_in": round(s.page_width / 914400, 3),
            "page_height_in": round(s.page_height / 914400, 3),
            "top_margin_in": round(s.top_margin / 914400, 3),
            "bottom_margin_in": round(s.bottom_margin / 914400, 3),
            "left_margin_in": round(s.left_margin / 914400, 3),
            "right_margin_in": round(s.right_margin / 914400, 3),
        })
    return {
        "path": str(path),
        "sha256": hash_file(path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "paragraph_styles": dict(styles.most_common()),
        "explicit_fonts_by_characters": dict(fonts.most_common()),
        "explicit_sizes_by_characters": {str(k): v for k, v in sizes.most_common()},
        "manual_bullet_paragraphs": manual_bullets,
        "section_layout": sections,
        "content_signature": content_signature(path),
    }


def validate_docx(path: str | Path, required_font: str = "Arial") -> dict[str, Any]:
    path = Path(path)
    doc = Document(path)
    findings: list[Finding] = []

    # Explicit non-standard fonts.
    for pi, p in enumerate(doc.paragraphs, 1):
        for ri, r in enumerate(p.runs, 1):
            if r.text.strip() and r.font.name and r.font.name.lower() != required_font.lower():
                findings.append(Finding("warning", "DOCX_FONT", f"Explicit font {r.font.name!r}; expected {required_font}.", f"paragraph {pi}, run {ri}"))
        text = p.text.lstrip()
        ppr = p._p.pPr
        has_num = bool(ppr is not None and ppr.numPr is not None)
        if text.startswith(("•", "", "●", "◦")) and not has_num:
            findings.append(Finding("warning", "DOCX_FAKE_BULLET", "Paragraph begins with a manual bullet character instead of Word list numbering.", f"paragraph {pi}"))

    # Table checks.
    for ti, table in enumerate(doc.tables, 1):
        if not table.rows:
            continue
        first_trpr = table.rows[0]._tr.trPr
        header_repeat = bool(first_trpr is not None and first_trpr.find(qn("w:tblHeader")) is not None)
        if len(table.rows) > 15 and not header_repeat:
            findings.append(Finding("warning", "DOCX_TABLE_HEADER", "Long table does not mark its first row as a repeating header.", f"table {ti}"))
        for ri, row in enumerate(table.rows, 1):
            trpr = row._tr.trPr
            no_split = bool(trpr is not None and trpr.find(qn("w:cantSplit")) is not None)
            if len(table.rows) <= 15 and not no_split:
                findings.append(Finding("info", "DOCX_ROW_SPLIT", "Table row may split across pages.", f"table {ti}, row {ri}"))
                break

    if not doc.paragraphs and not doc.tables:
        findings.append(Finding("error", "DOCX_EMPTY", "Document contains no paragraphs or tables."))

    errors = sum(f.severity == "error" for f in findings)
    warnings = sum(f.severity == "warning" for f in findings)
    return {
        "path": str(path),
        "sha256": hash_file(path),
        "status": "FAIL" if errors else ("WARN" if warnings else "PASS"),
        "errors": errors,
        "warnings": warnings,
        "findings": [f.to_dict() for f in findings],
    }


def compare_docx(source: str | Path, candidate: str | Path) -> dict[str, Any]:
    sdoc, cdoc = Document(source), Document(candidate)
    sp, cp = _paragraph_texts(sdoc), _paragraph_texts(cdoc)
    st, ct = _table_texts(sdoc), _table_texts(cdoc)
    diffs = []
    maxp = max(len(sp), len(cp))
    for i in range(maxp):
        a = sp[i] if i < len(sp) else None
        b = cp[i] if i < len(cp) else None
        if a != b:
            diffs.append({"type": "paragraph", "index": i + 1, "source": a, "candidate": b})
            if len(diffs) >= 50:
                break
    if len(diffs) < 50 and st != ct:
        diffs.append({"type": "tables", "message": "Table text or table structure differs."})
    return {
        "source": str(source),
        "candidate": str(candidate),
        "source_sha256": hash_file(source),
        "candidate_sha256": hash_file(candidate),
        "content_equal": sp == cp and st == ct,
        "source_paragraphs": len(sp),
        "candidate_paragraphs": len(cp),
        "source_tables": len(st),
        "candidate_tables": len(ct),
        "differences": diffs,
    }
