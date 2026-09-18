from __future__ import annotations

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1F4E79"
DARK_NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "E7E6E6"
INPUT_YELLOW = "FFF2CC"
GREEN = "70AD47"
WHITE = "FFFFFF"
BLACK = "000000"
BLUE_INPUT = "0000FF"

THIN_GREY = Side(style="thin", color="D9D9D9")


def apply_default_font(ws, name: str = "Arial", size: int = 10) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            current = cell.font
            cell.font = Font(
                name=name,
                size=size,
                bold=current.bold,
                italic=current.italic,
                color=current.color,
                underline=current.underline,
            )


def style_section_bar(cells, fill: str = NAVY) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")


def style_input(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=INPUT_YELLOW)
    cell.font = Font(name="Arial", size=10, color=BLUE_INPUT)


def financial_number_format(decimals: int = 0) -> str:
    if decimals <= 0:
        return '#,##0;(#,##0);-'
    zeros = "0" * decimals
    return f'#,##0.{zeros};(#,##0.{zeros});-'


def percentage_format(decimals: int = 1) -> str:
    zeros = "0" * decimals
    return f'0.{zeros}%;(0.{zeros}%);-'


def set_view(ws, zoom: int = 95, show_gridlines: bool = False, freeze: str | None = None) -> None:
    ws.sheet_view.zoomScale = zoom
    ws.sheet_view.showGridLines = show_gridlines
    if freeze:
        ws.freeze_panes = freeze


def fit_columns(ws, min_width: float = 8, max_width: float = 45, padding: float = 2) -> None:
    for idx in range(1, ws.max_column + 1):
        max_len = 0
        for cell in ws[get_column_letter(idx)]:
            if cell.value is None:
                continue
            max_len = max(max_len, max(len(part) for part in str(cell.value).splitlines()))
        ws.column_dimensions[get_column_letter(idx)].width = max(min_width, min(max_width, max_len + padding))
