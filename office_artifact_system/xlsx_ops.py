from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any
import json
import re
import zipfile

from openpyxl import load_workbook

from .common import Finding, hash_file

ERROR_VALUES = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def _cell_payload(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out[c.coordinate] = c.value
    return out


def _ooxml_feature_preflight(path: Path) -> dict[str, Any]:
    """Inspect package parts before choosing an editor that may drop advanced objects."""
    features: dict[str, Any] = {
        "vba": False,
        "pivot_tables": False,
        "slicers": False,
        "connections": False,
        "queries": False,
        "active_x": False,
        "embeddings": False,
        "external_links": False,
        "custom_xml": False,
    }
    try:
        with zipfile.ZipFile(path) as zf:
            names = [n.lower() for n in zf.namelist()]
    except (zipfile.BadZipFile, OSError):
        return {
            "risk": "UNKNOWN",
            "recommended_editor": "excel-com",
            "features": features,
            "reason": "Workbook package could not be safely inspected as OOXML.",
        }

    for name in names:
        if name.endswith("vbaproject.bin"):
            features["vba"] = True
        if "/pivottables/" in name or "/pivotcache/" in name:
            features["pivot_tables"] = True
        if "/slicers/" in name or "slicercache" in name:
            features["slicers"] = True
        if name.endswith("/connections.xml"):
            features["connections"] = True
        if "/queries/" in name or "querytables" in name:
            features["queries"] = True
        if "/activex/" in name:
            features["active_x"] = True
        if "/embeddings/" in name:
            features["embeddings"] = True
        if "/externallinks/" in name:
            features["external_links"] = True
        if name.startswith("customxml/"):
            features["custom_xml"] = True

    high = any(features[k] for k in ("active_x", "embeddings", "slicers", "queries", "connections"))
    medium = any(features[k] for k in ("vba", "pivot_tables", "external_links", "custom_xml"))
    if high:
        risk = "HIGH"
        editor = "excel-com-or-surgical-ooxml"
        reason = "Advanced/native Excel objects detected; avoid routine openpyxl save cycles."
    elif medium:
        risk = "MEDIUM"
        editor = "excel-com-preferred"
        reason = "Workbook contains features that merit native Excel preservation."
    else:
        risk = "LOW"
        editor = "openpyxl-or-excel-com"
        reason = "No high-risk OOXML parts were detected by package preflight."
    return {"risk": risk, "recommended_editor": editor, "features": features, "reason": reason}


def _serializable(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _style_key(cell) -> tuple:
    font = cell.font
    fill = cell.fill
    border = cell.border
    alignment = cell.alignment
    return (
        cell.style_id,
        cell.number_format,
        font.name,
        font.sz,
        bool(font.bold),
        bool(font.italic),
        getattr(font.color, "type", None),
        _serializable(getattr(font.color, "rgb", None)),
        fill.fill_type,
        getattr(fill.fgColor, "type", None),
        _serializable(getattr(fill.fgColor, "rgb", None)),
        border.left.style,
        border.right.style,
        border.top.style,
        border.bottom.style,
        alignment.horizontal,
        alignment.vertical,
        bool(alignment.wrap_text),
        int(alignment.indent or 0),
    )


def _sheet_structure(ws) -> dict[str, Any]:
    row_dims = {
        str(i): {"height": d.height, "hidden": bool(d.hidden), "outline": int(d.outlineLevel or 0)}
        for i, d in ws.row_dimensions.items()
        if d.height is not None or d.hidden or d.outlineLevel
    }
    col_dims = {
        str(k): {"width": d.width, "hidden": bool(d.hidden), "outline": int(d.outlineLevel or 0)}
        for k, d in ws.column_dimensions.items()
        if d.width is not None or d.hidden or d.outlineLevel
    }
    tables = sorted(getattr(ws.tables, "keys", lambda: [])())
    chart_types = sorted(type(c).__name__ for c in getattr(ws, "_charts", []) or [])
    try:
        cf_count = len(ws.conditional_formatting)
    except Exception:
        cf_count = 0
    try:
        dv_count = len(ws.data_validations.dataValidation)
    except Exception:
        dv_count = 0
    return {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_ranges": sorted(str(r) for r in ws.merged_cells.ranges),
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "zoom": ws.sheet_view.zoomScale,
        "show_gridlines": bool(ws.sheet_view.showGridLines),
        "sheet_state": ws.sheet_state,
        "tab_color": _serializable(getattr(ws.sheet_properties.tabColor, "rgb", None)) if ws.sheet_properties.tabColor else None,
        "row_dimensions": row_dims,
        "column_dimensions": col_dims,
        "tables": tables,
        "chart_types": chart_types,
        "conditional_format_count": cf_count,
        "data_validation_count": dv_count,
        "print_area": str(ws.print_area) if ws.print_area else None,
        "print_title_rows": ws.print_title_rows,
        "print_title_cols": ws.print_title_cols,
        "page_orientation": ws.page_setup.orientation,
        "fit_to_width": ws.page_setup.fitToWidth,
        "fit_to_height": ws.page_setup.fitToHeight,
    }


def inspect_xlsx(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    sheets = []
    for ws in wb.worksheets:
        formula_count = 0
        populated = 0
        errors = []
        fonts = Counter()
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                populated += 1
                if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                    formula_count += 1
                if c.data_type == "e" or (isinstance(c.value, str) and c.value in ERROR_VALUES):
                    errors.append(c.coordinate)
                if c.font and c.font.name:
                    fonts[c.font.name] += 1
        sheet = _sheet_structure(ws)
        sheet.update({
            "title": ws.title,
            "populated_cells": populated,
            "formula_cells": formula_count,
            "error_cells": errors,
            "fonts": dict(fonts.most_common()),
        })
        sheets.append(sheet)

    links = []
    for link in getattr(wb, "_external_links", []) or []:
        links.append(str(getattr(link, "file_link", None) or link))
    return {
        "path": str(path),
        "sha256": hash_file(path),
        "sheet_names": wb.sheetnames,
        "defined_names": sorted(str(n) for n in wb.defined_names),
        "external_link_count": len(links),
        "external_links": links,
        "edit_preflight": _ooxml_feature_preflight(path),
        "sheets": sheets,
    }


def validate_xlsx(path: str | Path, required_font: str = "Arial") -> dict[str, Any]:
    path = Path(path)
    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    try:
        values_wb = load_workbook(path, data_only=True, read_only=False, keep_links=True)
    except Exception:
        values_wb = None

    findings: list[Finding] = []
    if not wb.sheetnames:
        findings.append(Finding("error", "XLSX_NO_SHEETS", "Workbook contains no worksheets."))

    preflight = _ooxml_feature_preflight(path)
    if preflight["risk"] == "HIGH":
        findings.append(Finding(
            "warning",
            "XLSX_ADVANCED_FEATURES",
            "Advanced/native Excel objects detected; use Excel COM or surgical OOXML for mutations.",
        ))

    if getattr(wb, "_external_links", None):
        findings.append(Finding(
            "warning", "XLSX_EXTERNAL_LINK",
            f"Workbook contains {len(wb._external_links)} external link record(s)."
        ))

    for ws in wb.worksheets:
        values_ws = values_wb[ws.title] if values_wb and ws.title in values_wb.sheetnames else None
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if c.data_type == "e" or (isinstance(c.value, str) and c.value in ERROR_VALUES):
                    findings.append(Finding(
                        "error", "XLSX_ERROR_CELL",
                        f"Spreadsheet error value {c.value!r}.", f"{ws.title}!{c.coordinate}"
                    ))
                if c.font and c.font.name and c.font.name.lower() != required_font.lower():
                    findings.append(Finding(
                        "warning", "XLSX_FONT",
                        f"Font {c.font.name!r}; expected {required_font}.", f"{ws.title}!{c.coordinate}"
                    ))
                if isinstance(c.value, str) and c.value.startswith("=") and "#REF!" in c.value.upper():
                    findings.append(Finding(
                        "error", "XLSX_BROKEN_FORMULA",
                        "Formula contains #REF!.", f"{ws.title}!{c.coordinate}"
                    ))
                if c.data_type == "f" and values_ws is not None:
                    cached = values_ws[c.coordinate]
                    if cached.data_type == "e" or (
                        isinstance(cached.value, str) and cached.value in ERROR_VALUES
                    ):
                        findings.append(Finding(
                            "error", "XLSX_CALCULATED_ERROR",
                            f"Calculated formula result is {cached.value!r}.",
                            f"{ws.title}!{c.coordinate}",
                        ))

        if ws.sheet_state == "visible" and ws.max_row > 5 and ws.sheet_view.zoomScale and not (70 <= ws.sheet_view.zoomScale <= 130):
            findings.append(Finding(
                "warning", "XLSX_ZOOM",
                f"Unusual worksheet zoom {ws.sheet_view.zoomScale}%.", ws.title
            ))
        for key, dim in ws.column_dimensions.items():
            if dim.width and dim.width > 60:
                findings.append(Finding(
                    "warning", "XLSX_WIDE_COLUMN",
                    f"Column width {dim.width} is unusually wide.", f"{ws.title}!{key}:{key}"
                ))

    if len(findings) > 500:
        counts = Counter((f.severity, f.code, f.message) for f in findings)
        compact = []
        seen = set()
        for f in findings:
            key = (f.severity, f.code, f.message)
            if key in seen:
                continue
            seen.add(key)
            n = counts[key]
            compact.append(Finding(
                f.severity, f.code, f"{f.message} Occurrences: {n}.", f.location
            ))
        findings = compact

    errors = sum(f.severity == "error" for f in findings)
    warnings = sum(f.severity == "warning" for f in findings)
    return {
        "path": str(path),
        "sha256": hash_file(path),
        "status": "FAIL" if errors else ("WARN" if warnings else "PASS"),
        "errors": errors,
        "warnings": warnings,
        "edit_preflight": preflight,
        "findings": [f.to_dict() for f in findings],
    }


def _append_diff(diffs: list[dict], diff: dict, limit: int = 250) -> bool:
    if len(diffs) >= limit:
        return False
    diffs.append(diff)
    return len(diffs) < limit


def compare_xlsx(source: str | Path, candidate: str | Path) -> dict[str, Any]:
    source = Path(source)
    candidate = Path(candidate)
    swb = load_workbook(source, data_only=False, read_only=False, keep_links=True)
    cwb = load_workbook(candidate, data_only=False, read_only=False, keep_links=True)
    diffs: list[dict] = []

    if swb.sheetnames != cwb.sheetnames:
        _append_diff(diffs, {"type": "sheet_names", "source": swb.sheetnames, "candidate": cwb.sheetnames})

    s_names = sorted(str(n) for n in swb.defined_names)
    c_names = sorted(str(n) for n in cwb.defined_names)
    if s_names != c_names:
        _append_diff(diffs, {"type": "defined_names", "source": s_names, "candidate": c_names})

    common = [s for s in swb.sheetnames if s in cwb.sheetnames]
    for name in common:
        sws, cws = swb[name], cwb[name]
        sp, cp = _cell_payload(sws), _cell_payload(cws)
        coords = sorted(set(sp) | set(cp))
        for coord in coords:
            if sp.get(coord) != cp.get(coord):
                if not _append_diff(diffs, {
                    "type": "cell", "sheet": name, "cell": coord,
                    "source": sp.get(coord), "candidate": cp.get(coord)
                }):
                    break
            sc, cc = sws[coord], cws[coord]
            if sc.value is not None or cc.value is not None:
                if _style_key(sc) != _style_key(cc):
                    if not _append_diff(diffs, {
                        "type": "style", "sheet": name, "cell": coord,
                        "source": _style_key(sc), "candidate": _style_key(cc)
                    }):
                        break
        if len(diffs) >= 250:
            break

        ss, cs = _sheet_structure(sws), _sheet_structure(cws)
        for key in sorted(set(ss) | set(cs)):
            if ss.get(key) != cs.get(key):
                if not _append_diff(diffs, {
                    "type": "sheet_structure", "sheet": name, "field": key,
                    "source": ss.get(key), "candidate": cs.get(key)
                }):
                    break
        if len(diffs) >= 250:
            break

    source_preflight = _ooxml_feature_preflight(source)
    candidate_preflight = _ooxml_feature_preflight(candidate)
    if source_preflight["features"] != candidate_preflight["features"]:
        _append_diff(diffs, {
            "type": "ooxml_features",
            "source": source_preflight["features"],
            "candidate": candidate_preflight["features"],
        })

    return {
        "source": str(source),
        "candidate": str(candidate),
        "source_sha256": hash_file(source),
        "candidate_sha256": hash_file(candidate),
        "content_equal": not diffs,
        "difference_count_returned": len(diffs),
        "differences_truncated": len(diffs) >= 250,
        "source_edit_preflight": source_preflight,
        "candidate_edit_preflight": candidate_preflight,
        "differences": diffs,
    }
