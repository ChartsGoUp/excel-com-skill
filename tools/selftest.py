#!/usr/bin/env python3
from pathlib import Path
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

from office_artifact_system.docx_ops import validate_docx, compare_docx
from office_artifact_system.xlsx_ops import validate_xlsx, compare_xlsx


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        d1 = td / "a.docx"; d2 = td / "b.docx"
        doc = Document(); p = doc.add_paragraph("Test body"); p.runs[0].font.name = "Arial"; p.runs[0].font.size = Pt(10); doc.save(d1)
        doc.save(d2)
        assert validate_docx(d1)["errors"] == 0
        assert compare_docx(d1, d2)["content_equal"]

        x1 = td / "a.xlsx"; x2 = td / "b.xlsx"
        wb = Workbook(); ws = wb.active; ws.title = "Model"; ws["A1"] = "Label"; ws["A1"].font = Font(name="Arial", size=10); ws["B1"] = 1; ws["B1"].font = Font(name="Arial", size=10); wb.save(x1); wb.save(x2)
        assert validate_xlsx(x1)["errors"] == 0
        assert compare_xlsx(x1, x2)["content_equal"]
    print("selftest: PASS")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
